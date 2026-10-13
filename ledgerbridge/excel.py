"""Bounded, non-executing XLSX intake for the review workflow."""
from __future__ import annotations

import io
import posixpath
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

from .core import FIELDS, MONEY_FIELDS, WorkflowError, money, normalize_fields

MAX_BYTES = 1_000_000
MAX_ROWS = 500
MAX_UNCOMPRESSED = 5_000_000
HEADERS = (FIELDS - MONEY_FIELDS) | {"subtotal", "tax", "total"}


def _text(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        raise ValueError("boolean values are not accepted")
    text = str(value).strip()
    if len(text) > 2000:
        raise ValueError("value exceeds 2,000 characters")
    return text or None


def preview_xlsx(content: bytes, source_name: str):
    """Return deterministic mapped rows and row-level errors; never writes data."""
    if not isinstance(content, bytes) or not 0 < len(content) <= MAX_BYTES:
        raise WorkflowError("Excel file must be between 1 byte and 1 MB.")
    if not isinstance(source_name, str) or not source_name.lower().endswith(".xlsx") or len(source_name) > 200:
        raise WorkflowError("Choose an .xlsx workbook with a safe file name.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            info = archive.infolist()
            names = [item.filename for item in info]
            normalized_names = [name.replace("\\", "/") for name in names]
            if len(info) > 200 or len(names) != len(set(names)) or len(normalized_names) != len(set(normalized_names)) or sum(item.file_size for item in info) > MAX_UNCOMPRESSED:
                raise ValueError
            for item in info:
                normalized = item.filename.replace("\\", "/")
                first = normalized.split("/", 1)[0]
                if normalized.startswith("/") or ".." in normalized.split("/") or (":" in first) or posixpath.normpath(normalized).startswith("../") or normalized.lower().startswith("xl/externallinks/"):
                    raise ValueError
                if normalized.lower().endswith(".rels"):
                    relationships = ET.fromstring(archive.read(item))
                    if any(value.lower() == "external" for element in relationships.iter() for key, value in element.attrib.items() if key.rsplit("}", 1)[-1].lower() == "targetmode"):
                        raise ValueError
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception:
        raise WorkflowError("Excel workbook is invalid, unsafe, or exceeds intake limits.") from None
    try:
        if len(workbook.worksheets) != 1:
            raise WorkflowError("Excel intake requires exactly one worksheet.")
        sheet = workbook.active
        raw_headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1), ())]
        headers = [_text(value) for value in raw_headers]
        if not headers or any(not value for value in headers) or len(headers) != len(set(headers)) or set(headers) - HEADERS or not {"vendor", "date", "total"} <= set(headers):
            raise WorkflowError("Excel needs vendor, date, total headers. Optional: document_number, subtotal, tax, account_code, department, notes. Unknown or duplicate headers are rejected.")
        rows = []
        for number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
            if number > MAX_ROWS + 1:
                raise WorkflowError("Excel import limit is 500 records.")
            values = [cell.value for cell in cells[:len(headers)]]
            errors = []
            if len(cells) > len(headers) and any(cell.value is not None for cell in cells[len(headers):]):
                errors.append("Too many columns.")
            if any(cell.data_type == "f" for cell in cells):
                errors.append("Formulas are not accepted; upload displayed values only.")
            row = dict(zip(headers, values))
            fields = {}
            try:
                for key in FIELDS - MONEY_FIELDS:
                    fields[key] = _text(row.get(key))
                for key in MONEY_FIELDS:
                    fields[key] = money(row.get(key.removesuffix("_cents")))
                normalize_fields(fields)
            except (ValueError, WorkflowError) as exc:
                errors.append(str(exc))
            for required in ("vendor", "date", "total_cents"):
                if fields.get(required) in (None, ""):
                    errors.append(f"Required field missing: {required.removesuffix('_cents').replace('_', ' ')}.")
            rows.append({"row": number, "fields": fields, "errors": errors})
        if not rows:
            raise WorkflowError("Excel worksheet has no data rows.")
        return rows
    finally:
        workbook.close()
